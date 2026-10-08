#!/usr/bin/env python
"""
auditoria_facturas.py -- Auditoria de las facturas de V5 contra sus PDF de ARCA (S881, idea de Carlos 07/10).

Pedido: buscar los PDF de las facturas emitidas, comprobar que el registro de V5 (CAE, tipo, totales, renglones) coincide con lo que
dice ARCA y que los renglones de los pedidos reflejan lo facturado. Cards #161, #163, #143, #144.

MODO UNICO HOY: AUDITAR -- SOLO LECTURA. Abre la base con mode=ro y no escribe nada en ella. Corregir (con backup + PIN) es otra fase,
que se decide mirando este informe. El CAE no se "rehace": lo emite ARCA; aca se vuelve a LEER del PDF y se contrasta con el registro.

Fuentes documentales (de mas a menos confiable):
  1. PDF guardados dentro de la base (ingesta_facturas_raw.pdf_bytes).
  2. Carpetas con PDF (--pdf, se puede repetir): descargas copiadas al Silo (FACTURAS_PDF\\OF, FACTURAS_PDF\\P).
  3. Opcional --arca-csv: "Mis Comprobantes Emitidos" de ARCA (lo baja Carlos con su clave): lista COMPLETA de lo emitido.
Solo cuentan los comprobantes con fecha de emision >= --desde (por defecto 2026-01-01), por la fecha del propio comprobante.

Uso:
  python scripts/auditoria_facturas.py --db <copia_de_P.db> --pdf "Q:\\...\\FACTURAS_PDF\\OF" [--pdf ...] [--arca-csv archivo.csv] --salida informe.xlsx
"""
import argparse
import csv
import datetime
import hashlib
import io
import os
import re
import sqlite3
import sys
import unicodedata
from collections import defaultdict

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)

import fitz  # noqa: E402
fitz.TOOLS.mupdf_display_errors(False)
from backend.remitos.pdf_parser import extract_text_from_pdf, parse_invoice_data  # noqa: E402
from backend.productos.normalizacion import normalizar_producto, similitud, UMBRAL_SUGERENCIA  # noqa: E402

TOL = 0.011  # un centavo de tolerancia en importes
TIPO_ARCA = {1: "FACTURA_A", 6: "FACTURA_B", 11: "FACTURA_C", 3: "NOTA_CREDITO_A", 8: "NOTA_CREDITO_B", 13: "NOTA_CREDITO_C",
             2: "NOTA_DEBITO_A", 7: "NOTA_DEBITO_B", 12: "NOTA_DEBITO_C"}


def norm(s):
    t = unicodedata.normalize("NFKD", str(s or "")).encode("ASCII", "ignore").decode("ASCII").upper()
    return re.sub(r"\s+", " ", t).strip()


def num(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return None


def parse_importe(s):
    """'1.108.360,00' / '1108360,00' / '1108360.00' / '1,108,360.00' -> float."""
    s = str(s or "").strip().replace("$", "").replace(" ", "")
    if not s:
        return None
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


# ------------------------------------------------------------------ lectura de PDF
def leer_pdf(nombre, contenido):
    """-> (doc|None, motivo). doc normalizado desde el parser de la ingesta (backend/remitos/pdf_parser.py)."""
    if not contenido:
        return None, "archivo vacio"
    try:
        texto, palabras = extract_text_from_pdf(contenido)
        d = parse_invoice_data(texto, palabras)
    except Exception as e:  # noqa: BLE001
        return None, f"no se pudo leer: {type(e).__name__}"
    f = d.get("factura") or {}
    if not f.get("cae") or not f.get("numero"):
        return None, "sin CAE o sin numero (no es una factura de ARCA, o es una etiqueta)"
    try:
        pv, nro = (int(x) for x in str(f["numero"]).split("-"))
    except ValueError:
        return None, "numero ilegible"
    items = [{"descripcion": i.get("descripcion") or "", "cantidad": num(i.get("cantidad")) or 0.0,
              "precio": num(i.get("precio_unitario")) or 0.0, "subtotal": num(i.get("subtotal")) or 0.0,
              "alicuota": num(i.get("alicuota_iva")) if i.get("alicuota_iva") is not None else 21.0} for i in d.get("items") or []]
    neto, total = num(f.get("total_neto")), num(f.get("total_final"))
    suma = round(sum(i["subtotal"] for i in items), 2)
    avisos = []
    if f.get("audit_status") and f["audit_status"] != "VERIFICADO_OK":
        avisos.append(f"parser: {f['audit_status']}")
    if neto is not None and items and abs(suma - neto) > 0.5:
        avisos.append(f"los renglones suman {suma:.2f} y el neto del PDF es {neto:.2f}")
    if not items:
        avisos.append("no se leyeron renglones")
    return {
        "tipo": f.get("tipo_comprobante") or "", "pv": pv, "numero": nro, "fecha": f.get("fecha_emision") or "",
        "cae": str(f["cae"]), "vto": f.get("vto_cae") or "", "cuit": re.sub(r"\D", "", str((d.get("cliente") or {}).get("cuit") or "")),
        "receptor": (d.get("cliente") or {}).get("razon_social") or "", "neto": neto, "total": total, "items": items,
        "confianza": "OK" if not avisos else "REVISAR", "avisos": avisos,
    }, ""


def recolectar_pdf(db, carpetas, desde):
    """Une todas las fuentes por (pv, numero, cae). -> (docs: dict clave -> doc con 'fuentes', ignorados: list)."""
    docs, ignorados = {}, []

    def agregar(nombre, origen, contenido):
        doc, motivo = leer_pdf(nombre, contenido)
        if doc is None:
            ignorados.append((origen, nombre, motivo))
            return
        if doc["fecha"] and doc["fecha"] < desde:
            ignorados.append((origen, nombre, f"anterior a {desde} (emitida {doc['fecha']})"))
            return
        clave = (doc["pv"], doc["numero"], doc["cae"])
        sha = hashlib.sha1(contenido).hexdigest()[:8]
        if clave in docs:
            docs[clave]["fuentes"].append(f"{origen}: {nombre}")
            if docs[clave]["sha"] != sha:
                docs[clave]["variantes"] += 1
        else:
            doc["fuentes"] = [f"{origen}: {nombre}"]
            doc["sha"] = sha
            doc["variantes"] = 0
            docs[clave] = doc

    for rid, nombre, pdf in db.execute("select id, filename, pdf_bytes from ingesta_facturas_raw where pdf_bytes is not null"):
        agregar(nombre or f"raw {rid}", "base", pdf)
    for carpeta in carpetas:
        for raiz, _, archivos in os.walk(carpeta):
            for fn in sorted(archivos):
                if fn.lower().endswith(".pdf"):
                    with open(os.path.join(raiz, fn), "rb") as fh:
                        agregar(fn, os.path.basename(carpeta.rstrip("\\/")) or carpeta, fh.read())
    return docs, ignorados


# ------------------------------------------------------------------ lectura de V5
def leer_v5(db, desde):
    """-> (facturas por (pv, numero, cae), duplicadas: clave -> lista de registros). Una factura registrada DOS veces en V5 queda en `duplicadas`."""
    facturas, duplicadas = {}, defaultdict(list)
    for r in db.execute("""select f.id, f.tipo_comprobante, f.estado, f.punto_venta, f.numero_comprobante, f.fecha_emision, f.cae, f.cae_vencimiento,
                                  f.neto_gravado, f.iva_21, f.total, f.pedido_id, c.razon_social, c.cuit
                           from facturas f left join clientes c on c.id = f.cliente_id
                           where f.cae is not null and trim(f.cae) != ''"""):
        fid = r[0]
        items = [{"descripcion": d or "", "cantidad": num(q) or 0.0, "precio": num(p) or 0.0, "subtotal": num(s) or 0.0, "pedido_item_id": pi}
                 for d, q, p, s, pi in db.execute(
                     "select descripcion, cantidad, precio_unitario_neto, subtotal_neto, pedido_item_id from facturas_items where factura_id=? order by id", (fid,))]
        fecha = str(r[5] or "")[:10]
        if fecha and fecha < desde:
            continue
        clave_v5 = (int(r[3] or 0), int(r[4] or 0), str(r[6]).strip())
        registro = {
            "id": fid, "tipo": r[1], "estado": r[2], "pv": r[3], "numero": r[4], "fecha": fecha, "cae": str(r[6]).strip(), "vto": str(r[7] or "")[:10],
            "neto": num(r[8]), "iva": num(r[9]), "total": num(r[10]), "pedido_id": r[11], "receptor": r[12] or "",
            "cuit": re.sub(r"\D", "", str(r[13] or "")), "items": items}
        if clave_v5 in facturas:
            if clave_v5 not in duplicadas:   # (no usar `+=` sobre el defaultdict: crea la clave antes de evaluar la condicion)
                duplicadas[clave_v5].append(facturas[clave_v5])
            duplicadas[clave_v5].append(registro)
        else:
            facturas[clave_v5] = registro
    return facturas, dict(duplicadas)


def comparar_items(items_pdf, items_v5):
    """-> lista de diferencias. Cantidad y precio mandan (severidad CRITICA); el texto solo avisa."""
    dif = []
    a = sorted((round(i["cantidad"], 3), round(i["precio"], 2)) for i in items_pdf)
    b = sorted((round(i["cantidad"], 3), round(i["precio"], 2)) for i in items_v5)
    if a != b:
        dif.append(("CRITICA", "renglones",
                    "PDF: " + "; ".join(f"{c:g} x {p:g}" for c, p in a) + " | V5: " + "; ".join(f"{c:g} x {p:g}" for c, p in b)))
    else:
        ta = sorted(norm(i["descripcion"]) for i in items_pdf)
        tb = sorted(norm(i["descripcion"]) for i in items_v5)
        if ta != tb:
            dif.append(("TEXTO", "descripciones", "mismas cantidades y precios, distinto texto (el espejo copia el nombre del producto del pedido)"))
    return dif


def auditar_factura(pdf, v5):
    """Compara un documento de ARCA con el registro de V5. -> lista de (severidad, campo, detalle)."""
    dif = []
    if pdf["tipo"] and v5["tipo"] != pdf["tipo"]:
        dif.append(("TIPO", "tipo", f"ARCA: {pdf['tipo']} | V5: {v5['tipo']}"))
    if pdf["fecha"] and v5["fecha"] != pdf["fecha"]:
        dif.append(("MEDIA", "fecha", f"ARCA: {pdf['fecha']} | V5: {v5['fecha']}"))
    if pdf["vto"] and v5["vto"] and v5["vto"] != pdf["vto"]:
        dif.append(("MEDIA", "vencimiento CAE", f"ARCA: {pdf['vto']} | V5: {v5['vto']}"))
    if pdf["total"] is not None and (v5["total"] is None or abs(pdf["total"] - v5["total"]) > TOL):
        dif.append(("CRITICA", "total", f"ARCA: {pdf['total']:.2f} | V5: {v5['total'] if v5['total'] is None else format(v5['total'], '.2f')}"))
    if pdf["neto"] is not None and (v5["neto"] is None or abs(pdf["neto"] - v5["neto"]) > TOL):
        dif.append(("MEDIA", "neto", f"ARCA: {pdf['neto']:.2f} | V5: {v5['neto'] if v5['neto'] is None else format(v5['neto'], '.2f')}"))
    if pdf["cuit"] and v5["cuit"] and pdf["cuit"] != v5["cuit"]:
        dif.append(("CRITICA", "cliente", f"CUIT ARCA {pdf['cuit']} ({pdf['receptor']}) | V5 {v5['cuit']} ({v5['receptor']})"))
    dif += comparar_items(pdf["items"], v5["items"])
    return dif


# ------------------------------------------------------------------ candidatos de pedido para una factura que V5 no tiene
def pedidos_candidatos(db, pdf, dias=45):
    """Para una factura emitida que V5 no registra: pedidos del MISMO cliente (por CUIT), cerca en fecha y/o con el mismo total.
    -> texto corto con hasta 3 candidatos, o '' si no hay ninguno."""
    if not pdf["cuit"]:
        return ""
    clientes = [r[0] for r in db.execute("select id from clientes where replace(replace(coalesce(cuit,''),'-',''),' ','') = ?", (pdf["cuit"],))]
    if not clientes:
        return "(el cliente no existe en V5)"
    try:
        f0 = datetime.date.fromisoformat(pdf["fecha"])
    except ValueError:
        f0 = None
    cand = []
    for cid in clientes:
        for pid, fecha, total, estado in db.execute("select id, fecha, total, estado from pedidos where cliente_id=? and estado != 'ANULADO'", (cid,)):
            try:
                d = abs((datetime.date.fromisoformat(str(fecha)[:10]) - f0).days) if f0 else None
            except ValueError:
                d = None
            igual = pdf["total"] is not None and total is not None and abs(float(total) - pdf["total"]) <= 1.0
            if igual or (d is not None and d <= dias):
                ya = [r[0] for r in db.execute("select numero_comprobante from facturas where pedido_id=? and cae is not null and trim(cae)!=''", (str(pid),))]
                cand.append((0 if igual else 1, d if d is not None else 999, pid, fecha, total, estado, ya))
    cand.sort()
    return " || ".join(f"pedido #{c[2]} ({str(c[3])[:10]}, ${c[4]:,.0f}, {c[5]}){' TOTAL IGUAL' if c[0] == 0 else ''}{' ya tiene factura ' + ','.join(map(str, c[6])) if c[6] else ''}" for c in cand[:3])


# ------------------------------------------------------------------ pedido vs factura
def auditar_pedidos(db, docs_cruzados):
    """Por pedido con factura(s) auditadas: lo pedido vs lo facturado (segun ARCA), renglon por renglon.
    Clasificacion (la de la doctrina del Circuito PR): IGUAL / FACTURADO_DE_MAS (clase B) / FACTURADO_DE_MENOS / SIN_RENGLON_EN_PEDIDO."""
    por_pedido = defaultdict(list)
    for v5, pdf in docs_cruzados:
        if v5["pedido_id"] is not None:
            por_pedido[v5["pedido_id"]].append((v5, pdf))
    filas = []
    for pid, pares in sorted(por_pedido.items(), key=lambda kv: int(kv[0]) if str(kv[0]).isdigit() else 0):
        ped = db.execute("select id, estado from pedidos where id=?", (pid,)).fetchone()
        if not ped:
            for v5, pdf in pares:
                filas.append((pid, f"{pdf['pv']:05d}-{pdf['numero']:08d}", "(el pedido ya no existe)", None, sum(i["cantidad"] for i in pdf["items"]), "PEDIDO_INEXISTENTE", ""))
            continue
        lineas = [{"id": i, "nombre": n, "cant": q or 0.0, "precio": p or 0.0, "desc": di or 0.0, "fact": 0.0}
                  for i, n, q, p, di in db.execute(
                      """select pi.id, pr.nombre, pi.cantidad, pi.precio_unitario, pi.descuento_importe from pedidos_items pi
                         join productos pr on pr.id = pi.producto_id where pi.pedido_id=? order by pi.id""", (pid,))]
        sin_par = []
        for v5, pdf in pares:
            for it in pdf["items"]:
                mejor, mejor_s = None, 0.0
                for ln in lineas:
                    s = 1.0 if normalizar_producto(ln["nombre"]) == normalizar_producto(it["descripcion"]) else similitud(ln["nombre"], it["descripcion"])
                    if s > mejor_s:
                        mejor, mejor_s = ln, s
                if mejor is not None and mejor_s >= UMBRAL_SUGERENCIA:
                    mejor["fact"] += it["cantidad"]
                else:
                    sin_par.append((pdf, it))
        facturas_txt = ", ".join(sorted({f"{p['pv']:05d}-{p['numero']:08d}" for _, p in pares}))
        for ln in lineas:
            if abs(ln["fact"] - ln["cant"]) <= 0.001:
                estado = "IGUAL"
            elif ln["fact"] > ln["cant"]:
                estado = "FACTURADO_DE_MAS"
            elif ln["fact"] == 0:
                estado = "NO_FACTURADO_EN_ESTAS_FACTURAS"
            else:
                estado = "FACTURADO_DE_MENOS"
            filas.append((pid, facturas_txt, ln["nombre"], ln["cant"], ln["fact"], estado, ""))
        for pdf, it in sin_par:
            filas.append((pid, f"{pdf['pv']:05d}-{pdf['numero']:08d}", it["descripcion"], None, it["cantidad"], "SIN_RENGLON_EN_PEDIDO", "renglon de la factura que ningun renglon del pedido explica"))
    return filas


# ------------------------------------------------------------------ correcciones (fase APLICAR, como migracion con datos)
def desglose_iva(items, total_pdf):
    """Neto gravado, exento e IVA por alicuota calculados con los renglones. -> (dict, coincide_con_el_total_del_pdf)."""
    neto = sum(i["subtotal"] for i in items if i["alicuota"] > 0)
    exento = sum(i["subtotal"] for i in items if i["alicuota"] == 0)
    iva21 = sum(i["subtotal"] for i in items if abs(i["alicuota"] - 21.0) < 0.01) * 0.21
    iva105 = sum(i["subtotal"] for i in items if abs(i["alicuota"] - 10.5) < 0.01) * 0.105
    otras = [i for i in items if i["alicuota"] not in (0.0, 21.0, 10.5)]
    d = {"neto_gravado": round(neto, 2), "exento": round(exento, 2), "iva_21": round(iva21, 2), "iva_105": round(iva105, 2)}
    total = round(d["neto_gravado"] + d["exento"] + d["iva_21"] + d["iva_105"], 2)
    d["total"] = total
    return d, (not otras and total_pdf is not None and abs(total - total_pdf) <= 0.05)


def enlazar_renglones(db, pedido_id, items_pdf):
    """Empareja cada renglon del PDF con UN renglon del pedido (mismo producto: clave canonica o parecido con medidas exactas)."""
    if pedido_id is None or not str(pedido_id).isdigit():
        return [None] * len(items_pdf)
    lineas = [(i, n) for i, n in db.execute(
        "select pi.id, pr.nombre from pedidos_items pi join productos pr on pr.id = pi.producto_id where pi.pedido_id=? order by pi.id", (int(pedido_id),))]
    usados, salida = set(), []
    for it in items_pdf:
        mejor, mejor_s = None, 0.0
        for pid, nombre in lineas:
            if pid in usados:
                continue
            s = 1.0 if normalizar_producto(nombre) == normalizar_producto(it["descripcion"]) else similitud(nombre, it["descripcion"])
            if s > mejor_s:
                mejor, mejor_s = pid, s
        if mejor is not None and mejor_s >= UMBRAL_SUGERENCIA:
            usados.add(mejor)
            salida.append(mejor)
        else:
            salida.append(None)
    return salida


def construir_correccion(db, pdf, v):
    """Lo que habria que escribir en V5 para que la factura coincida con ARCA. Si algo no cierra, solo cabecera (tipo, fecha, vencimiento) o nada."""
    despues = {"tipo": pdf["tipo"], "fecha": pdf["fecha"], "vto": pdf["vto"]}
    nota = []
    desg, cierra = desglose_iva(pdf["items"], pdf["total"]) if pdf["items"] else ({}, False)
    items = None
    if cierra:
        despues.update(desg)
        if pdf["confianza"] == "OK":
            enlaces = enlazar_renglones(db, v["pedido_id"], pdf["items"])
            items = [{"descripcion": i["descripcion"], "cantidad": i["cantidad"], "precio": i["precio"], "alicuota": i["alicuota"],
                      "subtotal": i["subtotal"], "pedido_item_id": e} for i, e in zip(pdf["items"], enlaces)]
        else:
            nota.append("totales corregidos; renglones NO (lectura del PDF dudosa: " + "; ".join(pdf["avisos"]) + ")")
    elif pdf["neto"] is not None and pdf["total"] is not None and abs(pdf["neto"] * 1.21 - pdf["total"]) <= 0.05:
        # Los renglones no cierran (tipicamente por la bonificacion, Card #144) pero el encabezado impreso del PDF si es coherente:
        # neto + 21 % = total. Se toman los totales impresos; los renglones quedan como estan, a revisar a mano.
        despues.update({"neto_gravado": round(pdf["neto"], 2), "exento": 0.0, "iva_21": round(pdf["total"] - pdf["neto"], 2), "iva_105": 0.0, "total": round(pdf["total"], 2)})
        nota.append("totales tomados del encabezado impreso del PDF (neto + 21 % = total); renglones NO corregidos: los renglones leidos no cierran, revisar a mano (Card #144)")
    else:
        nota.append("totales y renglones NO corregidos: lo leido del PDF no cierra con el total impreso" if pdf["items"] else "sin renglones legibles: solo cabecera")
    return {"pv": pdf["pv"], "numero": pdf["numero"], "cae": pdf["cae"], "fecha_emision_arca": pdf["fecha"],
            "antes": {"tipo": v["tipo"], "total": v["total"], "neto": v["neto"], "n_items": len(v["items"]),
                      "suma_items": round(sum(i["subtotal"] for i in v["items"]), 2)},
            "despues": despues, "items": items, "observacion": " | ".join(nota)}


# ------------------------------------------------------------------ CSV de ARCA
def leer_csv_arca(ruta, desde):
    crudo = open(ruta, "rb").read()
    for enc in ("utf-8-sig", "latin-1"):
        try:
            texto = crudo.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    delim = ";" if texto.count(";") >= texto.count(",") else ","
    lector = csv.reader(io.StringIO(texto), delimiter=delim)
    filas = list(lector)
    if not filas:
        return []
    enc = [norm(c) for c in filas[0]]

    def col(*claves):
        for i, c in enumerate(enc):
            if all(k in c for k in claves):
                return i
        return None

    i_fecha, i_tipo, i_pv = col("FECHA"), col("TIPO"), col("PUNTO")
    i_nro = col("NUMERO", "DESDE") if col("NUMERO", "DESDE") is not None else col("NUMERO")
    i_cae = col("AUTORIZ") if col("AUTORIZ") is not None else col("CAE")
    # OJO: hay otra columna parecida, "Imp. Neto Gravado Total": el total del comprobante es la que se llama EXACTAMENTE "Imp. Total"
    i_total = next((i for i, c in enumerate(enc) if c == "IMP. TOTAL"), col("IMP", "TOTAL"))
    i_doc, i_den = col("NRO", "DOC"), col("DENOMINACION")
    salida = []
    for f in filas[1:]:
        if not f or len(f) <= max(x for x in (i_fecha, i_tipo, i_pv, i_nro, i_cae, i_total) if x is not None):
            continue
        fecha = f[i_fecha].strip()
        m = re.match(r"(\d{2})/(\d{2})/(\d{4})", fecha)
        fecha_iso = f"{m.group(3)}-{m.group(2)}-{m.group(1)}" if m else fecha[:10]
        if fecha_iso < desde:
            continue
        m_t = re.match(r"\s*(\d+)", f[i_tipo])
        try:
            salida.append({"fecha": fecha_iso, "tipo": TIPO_ARCA.get(int(m_t.group(1)), f[i_tipo].strip()) if m_t else f[i_tipo].strip(),
                           "pv": int(re.sub(r"\D", "", f[i_pv]) or 0), "numero": int(re.sub(r"\D", "", f[i_nro]) or 0),
                           "cae": re.sub(r"\D", "", f[i_cae]), "cuit": re.sub(r"\D", "", f[i_doc]) if i_doc is not None else "",
                           "receptor": f[i_den].strip() if i_den is not None else "", "total": parse_importe(f[i_total])})
        except (ValueError, IndexError):
            continue
    return salida


# ------------------------------------------------------------------ informe
def main():
    ap = argparse.ArgumentParser(description="Auditoria de facturas de V5 contra sus PDF de ARCA (solo lectura).")
    ap.add_argument("--db", required=True, help="copia de la base de P (se abre en solo lectura)")
    ap.add_argument("--pdf", action="append", default=[], help="carpeta con PDF de facturas (repetible)")
    ap.add_argument("--arca-csv", help="CSV 'Mis Comprobantes Emitidos' de ARCA")
    ap.add_argument("--desde", default="2026-01-01", help="fecha de emision minima (YYYY-MM-DD)")
    ap.add_argument("--salida", required=True, help="Excel de salida")
    ap.add_argument("--generar-correcciones", metavar="JSON", help="ademas del informe, escribe las correcciones (V5 -> lo que dice ARCA) para la migracion de datos")
    a = ap.parse_args()

    db = sqlite3.connect("file:" + os.path.abspath(a.db).replace("\\", "/") + "?mode=ro", uri=True)
    docs, ignorados = recolectar_pdf(db, a.pdf, a.desde)
    v5, duplicadas = leer_v5(db, a.desde)
    por_numero_v5 = defaultdict(list)
    for k, v in v5.items():
        por_numero_v5[(k[0], k[1])].append(v)

    filas, cruzados, resumen = [], [], defaultdict(int)
    usados_v5 = set()
    for clave, pdf in sorted(docs.items(), key=lambda kv: (kv[0][0], kv[0][1])):
        pv, nro, cae = clave
        if clave in duplicadas:
            regs = duplicadas[clave]
            resumen["DUPLICADA_EN_V5"] += 1
            usados_v5.add(clave)
            filas.append(("DUPLICADA_EN_V5 (decide Carlos)", f"{pv:05d}-{nro:08d}", pdf["tipo"], pdf["fecha"], pdf["cae"], pdf["receptor"], pdf["total"], "", "", None, None,
                          "V5 la registra " + str(len(regs)) + " veces: " + " || ".join(f"{r['tipo']}/{r['estado']} pedido {r['pedido_id']} total {r['total']}" for r in regs)
                          + " | no se corrige ninguna: hay que decidir cual queda", pdf["confianza"], " | ".join(pdf["fuentes"][:3])))
            continue
        v = v5.get(clave)
        if v is None:
            posibles = por_numero_v5.get((pv, nro), [])
            if posibles:
                v = posibles[0]
                dif = [("CRITICA", "CAE", f"el PDF trae CAE {cae} y V5 tiene {v['cae']} para el mismo numero")]
            else:
                dif = None
        else:
            dif = auditar_factura(pdf, v)
        if v is not None:
            usados_v5.add((v["pv"], v["numero"], v["cae"]))
        extra = ""
        if dif is None:
            resultado = "SOLO_PDF (emitida y V5 no la tiene)"
            if pdf["fecha"] and v5:
                f_min = min(x["fecha"] for x in v5.values() if x["fecha"])
                extra = ("pedidos posibles en V5: " + (pedidos_candidatos(db, pdf) or "ninguno")) if pdf["fecha"] >= f_min else "anterior al primer comprobante de V5"
        elif not dif:
            resultado = "OK"
        else:
            sev = {d[0] for d in dif}
            resultado = "DIFERENCIAS" + (" (CRITICAS)" if "CRITICA" in sev else "")
        if v is not None and dif is not None:
            cruzados.append((v, pdf))
        resumen[resultado.split(" ")[0]] += 1
        filas.append((resultado, f"{pv:05d}-{nro:08d}", pdf["tipo"], pdf["fecha"], pdf["cae"], pdf["receptor"], pdf["total"],
                      v["estado"] if v else "", v["tipo"] if v else "", v["total"] if v else None, v["pedido_id"] if v else None,
                      (extra + (" || " if extra and dif else "")) + " || ".join(f"[{s}] {c}: {d}" for s, c, d in (dif or [])), pdf["confianza"] + (": " + "; ".join(pdf["avisos"]) if pdf["avisos"] else ""),
                      " | ".join(pdf["fuentes"][:3])))
    solo_v5 = []
    for clave, v in sorted(v5.items(), key=lambda kv: (kv[0][0], kv[0][1])):
        if (v["pv"], v["numero"], v["cae"]) not in usados_v5:
            resumen["SOLO_V5"] += 1
            solo_v5.append((f"{v['pv']:05d}-{v['numero']:08d}", v["tipo"], v["estado"], v["fecha"], v["cae"], v["receptor"], v["total"], v["pedido_id"]))

    if a.generar_correcciones:
        import json
        correcciones = [construir_correccion(db, pdf, v) for v, pdf in cruzados if (v["pv"], v["numero"], v["cae"]) == (pdf["pv"], pdf["numero"], pdf["cae"])]
        with open(a.generar_correcciones, "w", encoding="utf-8") as fh:
            json.dump({"generado": datetime.datetime.now().isoformat(timespec="seconds"), "base": os.path.basename(a.db), "correcciones": correcciones}, fh, ensure_ascii=False, indent=1)
        completas = sum(1 for c in correcciones if c["items"] is not None)
        print(f"correcciones: {len(correcciones)} (con renglones: {completas}; solo cabecera/totales: {len(correcciones) - completas}) -> {a.generar_correcciones}")
        for c in correcciones:
            if c["observacion"]:
                print(f"   {c['pv']:05d}-{c['numero']:08d}: {c['observacion'][:140]}")

    filas_ped = auditar_pedidos(db, cruzados)
    arca = []
    if a.arca_csv:
        for r in leer_csv_arca(a.arca_csv, a.desde):
            en_v5 = v5.get((r["pv"], r["numero"], r["cae"]))
            en_pdf = docs.get((r["pv"], r["numero"], r["cae"]))
            if en_v5 and en_pdf:
                estado = "EN_V5_Y_CON_PDF"
            elif en_v5:
                estado = "EN_V5_SIN_PDF"
            elif en_pdf:
                estado = "NO_EN_V5_(CON_PDF)"
            else:
                estado = "NO_EN_V5_NI_PDF"
            tot = en_v5["total"] if en_v5 else (en_pdf["total"] if en_pdf else None)
            if tot is not None and r["total"] is not None and abs(tot - r["total"]) > TOL:
                estado += " | TOTAL DISTINTO"
            arca.append((estado, f"{r['pv']:05d}-{r['numero']:08d}", r["tipo"], r["fecha"], r["cae"], r["cuit"], r["receptor"], r["total"], tot))
            resumen["ARCA:" + estado.split(" ")[0]] += 1

    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()
    verde, amarillo, rojo, gris = (PatternFill("solid", fgColor=c) for c in ("E2F0D9", "FFF2CC", "F8CBAD", "EDEDED"))

    def hoja(titulo, cabecera, filas_h, color_col=None, anchos=None):
        ws = wb.create_sheet(titulo)
        ws.append(cabecera)
        for c in ws[1]:
            c.font = Font(bold=True); c.fill = gris; c.alignment = Alignment(wrap_text=True, vertical="top")
        for f in filas_h:
            ws.append(list(f))
            if color_col is not None:
                t = str(f[color_col])
                relleno = verde if t.startswith("OK") or t.startswith("IGUAL") or t.startswith("EN_V5_Y") else (rojo if ("CRITICA" in t or t.startswith("SOLO") or "NO_EN_V5" in t or "DE_MAS" in t or "INEXISTENTE" in t) else amarillo)
                ws.cell(row=ws.max_row, column=color_col + 1).fill = relleno
        ws.freeze_panes = "A2"
        for i, w in enumerate(anchos or [], start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        return ws

    ws0 = wb.active; ws0.title = "RESUMEN"
    ws0.append(["Auditoria de facturas -- " + datetime.date.today().isoformat() + " -- SOLO LECTURA"]); ws0["A1"].font = Font(bold=True, size=13)
    ws0.append([f"Base: {a.db}  |  Emitidas desde {a.desde}  |  PDF unicos leidos: {len(docs)}  |  PDF ignorados: {len(ignorados)}  |  Facturas con CAE en V5: {len(v5)}"])
    ws0.append([])
    for k, v in sorted(resumen.items()):
        ws0.append([k, v])
    ws0.append([])
    ws0.append(["Leyenda: OK = V5 coincide con el PDF de ARCA. DIFERENCIAS (CRITICAS) = importes, cantidades, precios, CAE o cliente distintos. TIPO = V5 guarda otro tipo de comprobante. TEXTO = solo cambia la descripcion."])
    ws0.append(["SOLO_PDF = emitida (hay PDF) y V5 no tiene registro. SOLO_V5 = V5 tiene CAE pero ningun PDF lo respalda. Nada de esto se corrigio: es un informe para decidir caso por caso."])
    ws0.column_dimensions["A"].width = 60
    hoja("FACTURAS", ["Resultado", "Numero", "Tipo (ARCA)", "Fecha", "CAE", "Receptor", "Total ARCA", "Estado V5", "Tipo V5", "Total V5", "Pedido",
                      "Diferencias", "Lectura del PDF", "Fuente"], filas, color_col=0, anchos=[26, 16, 16, 12, 17, 34, 14, 17, 15, 14, 9, 90, 40, 60])
    hoja("SOLO_V5", ["Numero", "Tipo V5", "Estado V5", "Fecha", "CAE", "Cliente", "Total V5", "Pedido"], solo_v5, anchos=[16, 15, 18, 12, 17, 36, 14, 9])
    hoja("PEDIDOS", ["Pedido", "Factura(s)", "Producto", "Cant. pedido", "Cant. facturada (ARCA)", "Clasificacion", "Nota"], filas_ped, color_col=5, anchos=[9, 30, 52, 13, 17, 34, 40])
    hoja("PDF_IGNORADOS", ["Fuente", "Archivo", "Motivo"], ignorados, anchos=[16, 70, 70])
    if a.arca_csv:
        hoja("ARCA_CSV", ["Estado", "Numero", "Tipo", "Fecha", "CAE", "CUIT receptor", "Receptor", "Total ARCA", "Total V5/PDF"], arca, color_col=0, anchos=[34, 16, 16, 12, 17, 15, 34, 14, 14])
    wb.save(a.salida)
    print(f"PDF unicos leidos: {len(docs)} | ignorados: {len(ignorados)} | facturas con CAE en V5 (>= {a.desde}): {len(v5)}")
    for k, v in sorted(resumen.items()):
        print(f"  {k}: {v}")
    print("Informe:", a.salida)


if __name__ == "__main__":
    main()
