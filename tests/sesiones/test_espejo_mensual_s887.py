r"""S887, pedido de Carlos (15/10): el Excel Espejo de Pedidos suma una pestaña por cada mes calendario (fecha, Nº de pedido, cliente, circuito, estado, importe), con chevrones de
filtro/orden, verde para Blanco y rosa chicle para Rosa, totales al pie, participacion por cliente con torta, y al terminar el año el archivo completo se guarda como historico.
Sobre una COPIA de la base de P del 14/10 17:12; no escribe en el Silo."""
import sys, os, shutil, sqlite3, importlib, tempfile, filecmp
from _esquema_p import copiar_p   # [S888] copia la base de P y le asegura la columna pedidos_items.kit_grupo (migrate_063)
from datetime import date
sys.path.insert(0, r"C:\dev\Sonido_Liquido_V5\scripts"); os.chdir(r"C:\dev\Sonido_Liquido_V5")
import openpyxl
checks = []
def check(d, c): checks.append((d, bool(c))); print(("[OK]  " if c else "[FAIL]"), d)
TMP = tempfile.mkdtemp(prefix="espejo_mensual_")
DB = os.path.join(TMP, "copia_P.db")
copiar_p(r"Q:\Mi unidad\V5_Silo_Claude\BACKUPS_DB\PRODUCCION_P\FIJAS_DE_PRUEBAS\PRODUCCION_P_al_20261014_1712.db", DB)
m = importlib.import_module("exportar_pedidos_excel")
NO_FISCAL, ANULADO = 4096, 1 << 35
con = sqlite3.connect(DB)
filas = con.execute("select p.id, substr(p.fecha,1,7), p.estado, p.flags_estado, p.total, c.razon_social from pedidos p join clientes c on c.id=p.cliente_id").fetchall()
con.close()
vigentes = [f for f in filas if not (f[3] & ANULADO) and f[2] != "ANULADO"]
anulados = {f[0] for f in filas if (f[3] & ANULADO) or f[2] == "ANULADO"}
check(f"partida: {len(filas)} pedidos, {len(anulados)} anulados, rosas {sum(1 for f in vigentes if f[3] & NO_FISCAL)}", len(filas) > 100 and len(anulados) > 5)

# ---- 1) las pestañas de un año con datos
wb = openpyxl.Workbook(); wb.active.title = "Pedidos TEST"
meses = m._hojas_mensuales(wb, DB, hoy=date(2026, 10, 15))
esperados = sorted({int(f[1][5:]) for f in vigentes if f[1][:4] == "2026"} | {10}, reverse=True)
check(f"una pestaña por cada mes con pedidos, el mes más nuevo primero, después de la hoja principal {wb.sheetnames}",
      meses == esperados and wb.sheetnames[0] == "Pedidos TEST" and wb.sheetnames[1] == "2026-10 Octubre" and wb.sheetnames[-1] == "2026-03 Marzo")
todos_ok = True; sumas_ok = True; color_ok = True; sin_anulados = True; filtro_ok = True; totales_ok = True
for mes in meses:
    ws = wb[f"2026-{mes:02d} {m.MESES[mes - 1]}"]
    delmes = [f for f in vigentes if f[1] == f"2026-{mes:02d}"]
    ult = 3 + len(delmes)
    datos = [[c.value for c in r] for r in ws.iter_rows(min_row=4, max_row=ult, max_col=6)]
    ids = [d[1] for d in datos]
    todos_ok &= sorted(ids) == sorted(f[0] for f in delmes) and [d[0] for d in datos] == sorted(d[0] for d in datos)
    sin_anulados &= not (set(ids) & anulados)
    sumas_ok &= abs(sum(d[5] for d in datos) - sum(f[4] or 0 for f in delmes)) < 0.01
    for r in ws.iter_rows(min_row=4, max_row=ult, max_col=6):
        rosa = r[3].value == "Rosa"
        esperado_rosa = bool(next(f for f in delmes if f[0] == r[1].value)[3] & NO_FISCAL)
        fill = r[0].fill.fgColor.rgb[-6:]
        color_ok &= (rosa == esperado_rosa) and all(c.fill.fgColor.rgb[-6:] == ("FF85C0" if rosa else "92D050") for c in r)
    filtro_ok &= ws.auto_filter.ref == f"A3:F{ult}" and ws.freeze_panes == "A4" and [c.value for c in ws[3][:6]] == ["Fecha", "Pedido Nº", "Cliente", "Circuito", "Estado", "Importe"]
    f_tot = {ws.cell(row=ult + 2 + k, column=5).value: ws.cell(row=ult + 2 + k, column=6).value for k in range(4)}
    totales_ok &= f_tot["TOTAL DEL MES"] == f"=SUBTOTAL(109,$F$4:$F${ult})" and f_tot["Pedidos"] == f"=SUBTOTAL(103,$B$4:$B${ult})" \
        and "SUMPRODUCT" in f_tot["Subtotal Blanco"] and '"Blanco"' in f_tot["Subtotal Blanco"] and '"Rosa"' in f_tot["Subtotal Rosa"]
check("cada mes lista exactamente sus pedidos vigentes (sin anulados), en orden de fecha", todos_ok and sin_anulados)
check("el importe de cada fila es el total guardado del pedido y la suma de cada mes coincide con la base", sumas_ok)
check("cada fila va en verde fuerte (92D050) si es Blanco y en rosa chicle (FF85C0) si es Rosa (bit NO_FISCAL_FORCE), la fila entera", color_ok)
check("chevrones: AutoFilter sobre la cabecera y todos los datos (A3:Fn), cabecera congelada, columnas Fecha / Pedido / Cliente / Circuito / Estado / Importe", filtro_ok)
check("totales al pie con SUBTOTAL (se recalculan con lo que deja el filtro): Pedidos, Subtotal Blanco, Subtotal Rosa y TOTAL DEL MES", totales_ok)

# ---- 2) participacion por cliente + torta
ws = wb["2026-10 Octubre"]
delmes = [f for f in vigentes if f[1] == "2026-10"]
ult = 3 + len(delmes)
check("el mes tiene UNA torta", len(ws._charts) == 1 and type(ws._charts[0]).__name__ == "PieChart")
cab = next(r[0].row for r in ws.iter_rows(min_row=ult, max_col=6) if r[2].value == "Cliente" and r[4].value == "Importe")
tabla = []
r = cab + 1
while ws.cell(row=r, column=3).value:
    tabla.append((ws.cell(row=r, column=3).value, ws.cell(row=r, column=5).value, ws.cell(row=r, column=6).value)); r += 1
por = {}
for f in delmes: por[f[5]] = por.get(f[5], 0) + (f[4] or 0)
check(f"participación por cliente: ordenada de mayor a menor, top {m.TOP_CLIENTES_TORTA} + «Otros», y suma el total del mes ({len(tabla)} filas)",
      len(tabla) <= m.TOP_CLIENTES_TORTA + 1 and [t[1] for t in tabla[:-1]] == sorted([t[1] for t in tabla[:-1]], reverse=True)
      and abs(sum(t[1] for t in tabla) - sum(por.values())) < 0.01 and tabla[0][0] == max(por, key=por.get) and tabla[-1][0] == "Otros")
check("cada fila trae el % del mes como fórmula y el gráfico muestra porcentajes", all(str(t[2]).startswith("=IF(SUM(") for t in tabla) and ws._charts[0].dataLabels.showPercent is True)

# ---- 3) mes sin pedidos y año nuevo
wb2 = openpyxl.Workbook(); wb2.active.title = "Pedidos TEST"
m._hojas_mensuales(wb2, DB, hoy=date(2026, 12, 1))
w = wb2["2026-12 Diciembre"]
check("el mes en curso aparece aunque no tenga pedidos: dice «Sin pedidos en el mes», sin filtro ni torta, con totales en 0",
      w["A4"].value == "Sin pedidos en el mes" and not w._charts and w.auto_filter.ref is None)
wb3 = openpyxl.Workbook(); wb3.active.title = "Pedidos TEST"
m._hojas_mensuales(wb3, DB, hoy=date(2027, 1, 5))
check(f"en el año nuevo solo hay pestañas del año nuevo (los pedidos de 2026 ya quedaron en el histórico) {wb3.sheetnames}", wb3.sheetnames == ["Pedidos TEST", "2027-01 Enero"])

# ---- 4) el espejo completo sigue igual: hoja principal primera y activa, _info con las pestañas
OUT = os.path.join(TMP, "espejo", "PEDIDOS_ESPEJO_TEST.xlsx"); os.makedirs(os.path.dirname(OUT))
m.DB_PATH, m.OUTPUT = DB, OUT
m.main()
wbm = openpyxl.load_workbook(OUT)
check(f"main(): la hoja de bloques sigue primera y activa, después las pestañas mensuales y al final _info {wbm.sheetnames[:2]}...{wbm.sheetnames[-1]}",
      wbm.sheetnames[0].startswith("Pedidos ") and wbm.active.title == wbm.sheetnames[0] and wbm.sheetnames[-1] == "_info" and any(n[:5] == f"{date.today().year}-" for n in wbm.sheetnames))
info = {r[0]: r[1] for r in wbm["_info"].iter_rows(values_only=True, max_col=2)}
check("_info registra la cantidad de pestañas mensuales", "Pestañas mensuales" in info)

# ---- 5) historico al terminar el año
def archivo_de(anio, nombre):
    ruta = os.path.join(TMP, nombre); wbx = openpyxl.Workbook(); wbx.active.title = "Pedidos X"
    wi = wbx.create_sheet("_info"); wi.append(["Generado", f"{anio}-12-31 23:30:00"]); wbx.save(ruta); return ruta
v = archivo_de(2026, "viejo.xlsx")
dest = m._archivar_si_cambio_de_anio(v, hoy=date(2027, 1, 2))
esperado = os.path.join(TMP, "HISTORICO", f"PEDIDOS_ESPEJO_{m.ENTORNO}_2026.xlsx")
check(f"el primer arranque del año nuevo guarda el espejo del año que termina en HISTORICO\\PEDIDOS_ESPEJO_{m.ENTORNO}_2026.xlsx (copia idéntica)", dest == esperado and filecmp.cmp(v, esperado, shallow=False))
wbv = openpyxl.load_workbook(v); wbv.create_sheet("cambio"); wbv.save(v)
check("un histórico ya guardado no se pisa", m._archivar_si_cambio_de_anio(v, hoy=date(2027, 2, 1)) is None and "cambio" not in openpyxl.load_workbook(esperado).sheetnames)
v2 = archivar = archivo_de(2026, "mismo_anio.xlsx")
check("en el mismo año no se archiva nada", m._archivar_si_cambio_de_anio(v2, hoy=date(2026, 10, 15)) is None)
check("sin archivo previo o sin _info no hace nada ni falla", m._archivar_si_cambio_de_anio(os.path.join(TMP, "no_existe.xlsx"), hoy=date(2027, 1, 2)) is None)

shutil.rmtree(TMP, ignore_errors=True)
falla = [d for d, ok in checks if not ok]
print(f"\nRESULTADO: {len(checks) - len(falla)}/{len(checks)} checks OK"); [print("FALLO:", d) for d in falla]
sys.exit(1 if falla else 0)
