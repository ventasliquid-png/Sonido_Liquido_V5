# backend/informes/export_utils.py
"""
Motor de exportación compartido del módulo Informes (S875,
DISENO_MODULO_INFORMES_S875_2026-09-26.md §2.2).

Recibe columnas + filas -- nunca un informe particular -- y devuelve CSV/Excel/PDF/TXT.
Se escribe una sola vez; los cinco informes de Nivel 1 y los de Nivel 2 lo reusan sin tocarlo.

`columnas`: lista de dicts {"key": "...", "label": "...", "width": opcional (Excel/PDF)}
`filas`: lista de dicts, cada uno con al menos las claves de `columnas` (ausentes -> vacío).
"""
import csv
import io
import unicodedata
from datetime import datetime
from typing import Callable, List, Dict, Any, Optional

from fastapi.responses import StreamingResponse


def _valor(fila: Dict[str, Any], key: str) -> Any:
    v = fila.get(key)
    return "" if v is None else v


def _nombre_archivo(titulo: str, extension: str) -> str:
    seguro = "".join(c if (c.isalnum() or c in "-_") else "_" for c in titulo)
    return f"{seguro}.{extension}"


def exportar_csv(titulo: str, columnas: List[Dict], filas: List[Dict]) -> StreamingResponse:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([c["label"] for c in columnas])
    for fila in filas:
        writer.writerow([_valor(fila, c["key"]) for c in columnas])
    contenido = "﻿" + buf.getvalue()  # BOM -- Excel abre UTF-8 con acentos sin romperse
    return StreamingResponse(
        iter([contenido]),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{_nombre_archivo(titulo, "csv")}"'},
    )


def exportar_excel(titulo: str, columnas: List[Dict], filas: List[Dict]) -> StreamingResponse:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = (titulo or "Informe")[:31]  # límite de Excel para nombres de hoja

    for col_idx, col in enumerate(columnas, 1):
        celda = ws.cell(row=1, column=col_idx, value=col["label"])
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1E3A8A")

    for row_idx, fila in enumerate(filas, 2):
        for col_idx, col in enumerate(columnas, 1):
            v = _valor(fila, col["key"])
            # Solo el vacío queda sin valor: un 0 (cantidad, $ facturado) es un dato, no una celda vacía.
            ws.cell(row=row_idx, column=col_idx, value=None if v == "" else v)

    for col_idx, col in enumerate(columnas, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = col.get("width", 18)
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{_nombre_archivo(titulo, "xlsx")}"'},
    )


def _texto_pdf(valor: Any) -> str:
    """Valor -> texto legible para el listado: fechas ISO a medianoche sin la hora, números con
    separador de miles y coma decimal (formato de la planilla), el resto tal cual."""
    if valor is None or valor == "":
        return ""
    if isinstance(valor, bool):
        return "Sí" if valor else "No"
    if isinstance(valor, (int, float)):
        if isinstance(valor, float) and not valor.is_integer():
            s = f"{valor:,.2f}"
        else:
            s = f"{int(valor):,}"
        return s.replace(",", "X").replace(".", ",").replace("X", ".")
    texto = str(valor)
    if len(texto) == 19 and texto[10] == "T" and texto.endswith("T00:00:00"):
        return texto[:10]
    return texto


def _generado() -> str:
    """Fecha y hora de generación: un archivo o una hoja impresa tiene que decir cuándo se sacó."""
    return datetime.now().strftime("generado %d/%m/%Y %H:%M")


def exportar_pdf(titulo: str, columnas: List[Dict], filas: List[Dict], subtitulo: Optional[str] = None) -> StreamingResponse:
    """Listado en PDF. Clase propia, NO la del motor de remitos: esa dibuja en cada página el marco y
    el rótulo ORIGINAL/DUPLICADO/TRIPLICADO del remito y deja el cursor al pie, con lo que cada celda
    disparaba un salto de página (un listado de 178 filas salió con 3751 páginas y ningún renglón)."""
    from fpdf import FPDF
    from backend.remitos.remito_engine import PDFRemito

    safe = PDFRemito._pdf_safe
    orientacion = "L" if len(columnas) > 5 else "P"

    class _PDFInforme(FPDF):
        def footer(self):
            self.set_y(-10)
            self.set_font("Helvetica", "", 7)
            self.set_text_color(120, 120, 120)
            self.cell(0, 5, safe(f"{titulo} - pagina {self.page_no()} de {{nb}}"), align="C")

    pdf = _PDFInforme(orientation=orientacion, unit="mm", format="A4")
    pdf.set_margins(8, 10, 8)
    pdf.set_auto_page_break(auto=True, margin=14)
    pdf.alias_nb_pages()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 9, safe(titulo), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 8)
    pdf.cell(0, 5, safe(f"{len(filas)} fila(s) - {_generado()}"), new_x="LMARGIN", new_y="NEXT")
    if subtitulo:
        pdf.set_font("Helvetica", "I", 7)
        pdf.multi_cell(0, 4, safe(subtitulo), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)

    # Ancho de cada columna = lo que pide su contenido real (medido con la fuente, cabecera incluida).
    # Si la suma no entra en la hoja se achican todas en proporción; si sobra, se reparte parejo.
    # (El "width" de las columnas está pensado para Excel, en caracteres: no sirve para medir PDF.)
    ancho_util = pdf.w - pdf.l_margin - pdf.r_margin
    tam = 7 if len(columnas) <= 11 else 6.5   # muchas columnas: letra más chica (la hoja ya va apaisada)
    pdf.set_font("Helvetica", "B", tam)
    natural = [pdf.get_string_width(safe(c["label"])) + 3 for c in columnas]
    pdf.set_font("Helvetica", "", tam)
    for fila in filas:
        for i, col in enumerate(columnas):
            natural[i] = max(natural[i], pdf.get_string_width(safe(_texto_pdf(_valor(fila, col["key"])))) + 3)
    # Tope por columna: un único texto larguísimo (un cliente con tres razones sociales, un fragmento de
    # nota) no puede acaparar la hoja y apretar al resto; ese caso puntual se recorta con "...".
    tope = ancho_util * 0.28
    natural = [min(n, tope) for n in natural]
    if sum(natural) <= ancho_util:
        extra = (ancho_util - sum(natural)) / len(natural)
        anchos = [n + extra for n in natural]
    else:
        # No entra: las columnas cortas conservan lo que necesitan y solo se achican las largas, en
        # proporción (achicar todas por igual recortaba hasta "Nota humana" o un número de factura).
        anchos = [0.0] * len(natural)
        pendientes = set(range(len(natural)))
        restante = ancho_util
        while pendientes:
            parte = restante / len(pendientes)
            cortas = [i for i in pendientes if natural[i] <= parte]
            if not cortas:
                break
            for i in cortas:
                anchos[i] = natural[i]
                restante -= natural[i]
                pendientes.discard(i)
        if pendientes:
            total_largas = sum(natural[i] for i in pendientes)
            for i in pendientes:
                anchos[i] = restante * natural[i] / total_largas
    alto = 5.5

    def _ajustar(texto: str, ancho: float) -> str:
        """Recorta con '...' lo que no entra en la celda (medido con la fuente, no por cantidad de letras)."""
        maximo = ancho - 2
        if pdf.get_string_width(texto) <= maximo:
            return texto
        while texto and pdf.get_string_width(texto + "...") > maximo:
            texto = texto[:-1]
        return texto + "..."

    def _header():
        pdf.set_font("Helvetica", "B", tam)
        pdf.set_fill_color(30, 58, 138)
        pdf.set_text_color(255, 255, 255)
        for col, ancho in zip(columnas, anchos):
            pdf.cell(ancho, 6.5, _ajustar(safe(col["label"]), ancho), border=1, fill=True)
        pdf.ln()
        pdf.set_text_color(0, 0, 0)
        pdf.set_font("Helvetica", "", tam)

    _header()
    for i, fila in enumerate(filas):
        if pdf.get_y() + alto > pdf.h - 14:
            pdf.add_page()
            _header()
        pdf.set_fill_color(241, 245, 249) if i % 2 else pdf.set_fill_color(255, 255, 255)
        for col, ancho in zip(columnas, anchos):
            texto = safe(_texto_pdf(_valor(fila, col["key"])))
            pdf.cell(ancho, alto, _ajustar(texto, ancho), border=1, fill=True)
        pdf.ln()

    salida = bytes(pdf.output())
    return StreamingResponse(
        io.BytesIO(salida),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{_nombre_archivo(titulo, "pdf")}"'},
    )


def exportar_txt(
    titulo: str,
    columnas: List[Dict],
    filas: List[Dict],
    formato_linea: Optional[Callable[[Dict], str]] = None,
    subtitulo: Optional[str] = None,
) -> StreamingResponse:
    """Lista para leer o pegar (WhatsApp/mail), no una tabla -- un resultado por
    párrafo, en texto corrido. `formato_linea` permite que un informe (ej. E) arme
    su propio párrafo legible en vez del volcado genérico clave: valor."""
    buf = io.StringIO()
    buf.write(f"{titulo}\n{'=' * len(titulo)}\n{len(filas)} fila(s) - {_generado()}\n")
    if subtitulo:
        buf.write(f"{subtitulo}\n")
    buf.write("\n")
    for fila in filas:
        if formato_linea:
            buf.write(formato_linea(fila) + "\n\n")
        else:
            partes = [f"{c['label']}: {_valor(fila, c['key'])}" for c in columnas]
            buf.write(" | ".join(partes) + "\n\n")
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="text/plain; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{_nombre_archivo(titulo, "txt")}"'},
    )


EXPORTADORES = {
    "csv": exportar_csv,
    "excel": exportar_excel,
    "pdf": exportar_pdf,
    "txt": exportar_txt,
}


def exportar(formato: str, titulo: str, columnas: List[Dict], filas: List[Dict], **kwargs) -> StreamingResponse:
    fn = EXPORTADORES.get(formato)
    if not fn:
        raise ValueError(f"Formato de exportación desconocido: {formato}")
    if formato == "txt":
        return fn(titulo, columnas, filas, formato_linea=kwargs.get("formato_linea"), subtitulo=kwargs.get("subtitulo"))
    if formato == "pdf":
        return fn(titulo, columnas, filas, subtitulo=kwargs.get("subtitulo"))
    return fn(titulo, columnas, filas)


def normalizar_texto(valor: Optional[str]) -> str:
    """[Informe E] Saca acentos para comparar sin distinguirlos (á->a, ñ->n, etc.)
    -- unicodedata, sin dependencia nueva. 'medico' debe encontrar 'médico'."""
    if not valor:
        return ""
    forma = unicodedata.normalize("NFKD", valor)
    return "".join(c for c in forma if not unicodedata.combining(c)).lower()
